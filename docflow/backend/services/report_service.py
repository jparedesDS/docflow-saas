import pandas as pd
import io
import shutil
import os
from typing import List, Dict, Any
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Border, Side, Alignment
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.formatting.rule import Rule, DataBarRule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter


# ─── Constantes de estilos (idénticas al DocuControl original) ───────────────

FILL_HEADER   = PatternFill("solid", fgColor="6678AF")   # azul grisáceo
FILL_ROW      = PatternFill("solid", fgColor="D4DCF4")   # azul muy claro (todas las hojas)
FONT_WHITE    = Font(color="FFFFFF", bold=True)
FONT_BLACK    = Font(color="000000")
THIN_BORDER   = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)
ALIGN_CENTER  = Alignment(horizontal="center", vertical="center")
ALIGN_VCE     = Alignment(vertical="center")

# Colores por estado en celda Estado
ESTADO_COLORS = {
    "Rechazado":   "FFA19A",
    "Com. Menores":"FFE5AD",
    "Com. Mayores":"DBB054",
    "Comentado":   "F79646",
    "Enviado":     "B1E1B9",
    "Sin Enviar":  "FFFFAB",
    "Información": "FFFF46",
    "HOLD":        "FF0909",
    "Aprobado":    "00D25F",
}

# Colores de texto por responsable
RESPONSABLE_COLORS = {
    "SS":    "262626",
    "CCH":   "00B0F0",
    "JM":    "2C8A6A",
    "JV":    "006B95",
    "EC":    "B31274",
    "ES":    "5A0DA0",
    "JP":    "00458F",
    "AC":    "3F0075",
    "LB":    "176DD1",
    "RM":    "228B22",
    "RP":    "1B365D",
    "EC/SS": "1F7A1F",
}

# Colores de pestaña (igual que el original)
TAB_COLORS = {
    "ALL DOC.":      "6678AF",
    "ENVIADOS":      "00D25F",
    "DEVOLUCIONES":  "FFA19A",
    "CRÍTICOS":      "DBB054",
    "CRÍTICOS +15d": "FF7F50",
    "SIN ENVIAR":    "FFFF66",
    "STATUS GLOBAL": "B1E1B9",
}

# Orden de columnas (igual que el original)
COLUMN_ORDER = [
    "Nº Pedido", "Responsable", "Nº Oferta", "Nº PO", "Cliente", "Material",
    "Fecha Pedido", "Fecha Prevista", "Nº Doc. Cliente", "Nº Doc. EIPSA",
    "Título", "Tipo Doc.", "Info/Review", "Repsonsable", "Días Envío",
    "Crítico", "Estado", "Notas", "Nº Revisión",
    "Fecha Env. Doc.", "Fecha Dev. Doc.", "Días Devolución",
    "Reclamaciones", "Seguimiento", "Historial Rev.",
]

# Columnas que se centran
COLS_CENTRAR = {
    "Responsable", "Repsonsable", "Info/Review", "Días Envío", "Crítico",
    "Estado", "Días Devolución", "Nº Revisión",
    "Aprobado", "Com. Mayores", "Com. Menores", "Enviado", "Rechazado",
    "Sin Enviar", "Total", "% Completado",
}

DATE_COLS = {"Fecha Pedido", "Fecha Prevista", "Fecha Env. Doc.", "Fecha Dev. Doc."}


class ReportService:

    # ──────────────────────────────────────────────────────────────────────────
    # Generador principal del Monitoring Excel
    # ──────────────────────────────────────────────────────────────────────────

    def generate_monitoring_excel(self, sections: Dict) -> bytes:
        """Genera Excel multi-hoja con exactamente el mismo estilo que el DocuControl original."""
        all_docs      = sections.get("all_docs") or (
            sections.get("enviados", []) +
            sections.get("devoluciones", []) +
            sections.get("sin_enviar", []) +
            [d for d in sections.get("criticos", [])
             if not any(d.get("Nº Doc. EIPSA") == x.get("Nº Doc. EIPSA")
                        for x in sections.get("enviados", []) + sections.get("devoluciones", []))]
        )
        enviados      = sections.get("enviados", [])
        devoluciones  = sections.get("devoluciones", [])
        sin_enviar    = sections.get("sin_enviar", [])

        # Split críticos: excluir Enviado, luego ≤15 días → CRÍTICOS, >15 días → CRÍTICOS +15d
        criticos_raw  = [d for d in sections.get("criticos", [])
                         if str(d.get("Estado", "") or "").strip().lower() != "enviado"]
        criticos      = [d for d in criticos_raw
                         if _dias_dev(d) is None or _dias_dev(d) <= 15]
        criticos_15d  = [d for d in criticos_raw
                         if _dias_dev(d) is not None and _dias_dev(d) > 15]

        # SIN ENVIAR: normalizar Estado vacío → "Sin Enviar"
        sin_enviar    = [dict(d, Estado="Sin Enviar") if not str(d.get("Estado") or "").strip() else d
                         for d in sin_enviar]

        # Añadir columna "Notas" a devoluciones
        devoluciones  = [_add_notas(d) for d in devoluciones]

        wb = Workbook()
        wb.remove(wb.active)

        sheet_order = [
            ("ALL DOC.",      all_docs,     False),
            ("ENVIADOS",      enviados,     True),
            ("DEVOLUCIONES",  devoluciones, True),
            ("CRÍTICOS",      criticos,     True),
            ("CRÍTICOS +15d", criticos_15d, True),
            ("SIN ENVIAR",    sin_enviar,   False),
        ]

        for sheet_name, data, has_dias_dev_format in sheet_order:
            ws = wb.create_sheet(title=sheet_name)
            ws.sheet_properties.tabColor = TAB_COLORS[sheet_name]

            cols = _get_cols(data, sheet_name)
            _write_header(ws, cols)
            _write_rows(ws, data, cols)
            _apply_cell_styles(ws, cols)

            # Formato condicional: fila completa rosa en DEVOLUCIONES si Días Dev > 15
            if sheet_name == "DEVOLUCIONES" and "Días Devolución" in cols:
                col_idx = cols.index("Días Devolución") + 1
                col_letter = get_column_letter(col_idx)
                max_col_letter = get_column_letter(len(cols))
                if ws.max_row >= 2:
                    diff_row = DifferentialStyle(
                        fill=PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                    )
                    formula = f"=${col_letter}2>15"
                    rule_row = Rule(type="expression", formula=[formula], dxf=diff_row, stopIfTrue=False)
                    ws.conditional_formatting.add(f"A2:{max_col_letter}{ws.max_row}", rule_row)

            _finalize_sheet(ws, cols)

        # Hoja STATUS GLOBAL
        self._build_status_global_sheet(wb, all_docs)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    def _build_status_global_sheet(self, wb: Workbook, all_docs: list):
        """Construye la hoja STATUS GLOBAL con gráfico de barras apiladas."""
        if not all_docs:
            return

        df = pd.DataFrame(all_docs)
        if "Nº Pedido" not in df.columns or "Estado" not in df.columns:
            return

        df["Estado"] = df["Estado"].fillna("Sin Enviar").astype(str)

        status_global = (
            df.groupby(["Nº Pedido", "Estado"])
            .size()
            .unstack(fill_value=0)
            .reset_index()
        )
        # Asegurar que existen todas las columnas de estado
        for col in ["Aprobado", "Com. Mayores", "Com. Menores", "Enviado", "Rechazado", "Sin Enviar"]:
            if col not in status_global.columns:
                status_global[col] = 0
        status_global["Total"] = status_global.iloc[:, 1:].sum(axis=1)
        status_global["% Completado"] = (
            status_global.get("Aprobado", 0) / status_global["Total"] * 100
        ).fillna(0).round(2)

        # Excluir pedidos ya completados al 100%
        status_global = status_global[status_global["% Completado"] != 100].copy()
        status_global.sort_values("Nº Pedido", ascending=False, inplace=True)

        sg_cols = ["Nº Pedido", "Aprobado", "Com. Mayores", "Com. Menores",
                   "Enviado", "Rechazado", "Sin Enviar", "Total", "% Completado"]
        # Solo columnas disponibles
        sg_cols = [c for c in sg_cols if c in status_global.columns]

        ws = wb.create_sheet(title="STATUS GLOBAL")
        ws.sheet_properties.tabColor = TAB_COLORS["STATUS GLOBAL"]

        _write_header(ws, sg_cols)

        fill_sg = PatternFill("solid", fgColor="D4DCF4")
        for _, row_data in status_global.iterrows():
            ws.append([row_data.get(c, 0) for c in sg_cols])

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.fill = fill_sg
                cell.border = THIN_BORDER
                cell.font = FONT_BLACK
                col_name = sg_cols[cell.column - 1] if cell.column <= len(sg_cols) else ""
                if col_name in COLS_CENTRAR:
                    cell.alignment = ALIGN_CENTER
                else:
                    cell.alignment = ALIGN_VCE
                # % Completado = 100 → rojo bold (aunque ya excluimos 100%)
                if col_name == "% Completado" and cell.value == 100:
                    cell.font = Font(color="FF0000", bold=True)

        # DataBar en columna % Completado
        pct_idx = sg_cols.index("% Completado") + 1 if "% Completado" in sg_cols else None
        if pct_idx and ws.max_row >= 2:
            pct_letter = get_column_letter(pct_idx)
            ws.conditional_formatting.add(
                f"{pct_letter}2:{pct_letter}{ws.max_row}",
                DataBarRule(start_type="percentile", start_value=0,
                            end_type="percentile", end_value=100,
                            color="4472C4", showValue="1"),
            )

        _finalize_sheet(ws, sg_cols)

        # Gráfico de barras apiladas
        if ws.max_row >= 3:
            estado_cols_chart = ["Aprobado", "Com. Mayores", "Com. Menores", "Enviado", "Rechazado", "Sin Enviar"]
            chart_col_indices = [sg_cols.index(c) + 1 for c in estado_cols_chart if c in sg_cols]
            if chart_col_indices:
                min_col_c = min(chart_col_indices)
                max_col_c = max(chart_col_indices)

                chart = BarChart()
                chart.type = "col"
                chart.title = "Estado por Pedido"
                chart.style = 12
                chart.grouping = "stacked"
                chart.overlap = 100
                chart.y_axis.title = "Nº Documentos"
                chart.x_axis.title = "Nº Pedido"

                data_ref = Reference(ws, min_col=min_col_c, max_col=max_col_c,
                                     min_row=1, max_row=ws.max_row)
                cats = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
                chart.add_data(data_ref, titles_from_data=True)
                chart.set_categories(cats)

                colores_chart = ["00B350", "C59B3F", "FFCF7F", "5566A0", "FF8273", "FFEF7F"]
                for idx, serie in enumerate(chart.series):
                    if idx < len(colores_chart):
                        serie.graphicalProperties = GraphicalProperties(
                            solidFill=colores_chart[idx]
                        )

                chart.height = 16
                chart.width = 29
                ws.add_chart(chart, "J3")

    # ──────────────────────────────────────────────────────────────────────────
    # Otros métodos
    # ──────────────────────────────────────────────────────────────────────────

    def copy_to_shared_drive(self, excel_bytes: bytes, shared_path: str) -> Dict:
        today = datetime.now().strftime("%Y-%m-%d")
        filename = f"Monitoring_Report_{today}.xlsx"
        full_path = os.path.join(shared_path, filename)
        try:
            os.makedirs(shared_path, exist_ok=True)
            with open(full_path, "wb") as f:
                f.write(excel_bytes)
            return {"ok": True, "path": full_path, "filename": filename}
        except Exception as e:
            return {"ok": False, "error": str(e)}

    def generate_excel_report(self, documents: List[Dict[str, Any]], filters: dict = None) -> bytes:
        df = pd.DataFrame(documents)
        if filters:
            for key, value in filters.items():
                if value and key in df.columns:
                    df = df[df[key].astype(str).str.contains(str(value), case=False, na=False)]
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Documentos")
            summary_data = self._build_summary(df)
            pd.DataFrame(summary_data).to_excel(writer, index=False, sheet_name="Resumen")
        return output.getvalue()

    def generate_summary(self, documents: List[Dict[str, Any]]) -> dict:
        df = pd.DataFrame(documents)
        if df.empty:
            return {"total": 0, "by_status": {}, "generated_at": datetime.now().isoformat()}
        status_col = next((c for c in df.columns if c.lower() in ["estado", "status"]), None)
        by_status = df[status_col].value_counts().to_dict() if status_col else {}
        return {
            "total": len(df),
            "by_status": {str(k): int(v) for k, v in by_status.items()},
            "columns": list(df.columns),
            "generated_at": datetime.now().isoformat(),
        }

    def _build_summary(self, df: pd.DataFrame) -> list:
        rows = [{"Metrica": "Total documentos", "Valor": len(df)}]
        for col in df.columns:
            if col.lower() in ["estado", "status"]:
                for status, count in df[col].value_counts().items():
                    rows.append({"Metrica": f"Estado: {status}", "Valor": count})
        return rows


# ─── Funciones auxiliares ────────────────────────────────────────────────────

def _dias_dev(doc: dict):
    """Devuelve Días Devolución como int o None."""
    v = doc.get("Días Devolución")
    try:
        return int(float(v)) if v not in (None, "", "nan") else None
    except (ValueError, TypeError):
        return None


def _add_notas(doc: dict) -> dict:
    """Añade columna Notas a un doc de devoluciones: 'Enviar antes del DD/MM/YYYY'."""
    d = dict(doc)
    fecha = d.get("Fecha Env. Doc.")
    if fecha:
        try:
            if isinstance(fecha, str):
                # Intentar parsear la fecha
                clean = fecha.split("T")[0] if "T" in fecha else fecha
                dt = datetime.strptime(clean, "%Y-%m-%d")
            else:
                dt = fecha
            limite = dt + timedelta(days=15)
            d["Notas"] = f"Enviar antes del {limite.strftime('%d/%m/%Y')}"
        except Exception:
            d["Notas"] = ""
    else:
        d["Notas"] = ""
    return d


def _get_cols(data: list, sheet_name: str) -> list:
    """Devuelve columnas en el orden canónico, solo las que existen en los datos."""
    if not data:
        return list(COLUMN_ORDER)
    available = set(data[0].keys())
    ordered = [c for c in COLUMN_ORDER if c in available]
    # Añadir columnas extra no contempladas en COLUMN_ORDER
    extras = [c for c in data[0].keys() if c not in ordered]
    return ordered + extras


def _fmt_value(val, col_name: str):
    """Convierte un valor al tipo adecuado para la celda Excel."""
    if val is None or val == "":
        return ""
    if col_name in DATE_COLS:
        if isinstance(val, datetime):
            return val
        try:
            clean = str(val).split("T")[0]
            return datetime.strptime(clean, "%Y-%m-%d")
        except Exception:
            return str(val)
    if col_name in {"Días Envío", "Días Devolución", "Nº Revisión"}:
        try:
            n = int(float(val))
            return n if str(val) not in ("", "nan") else ""
        except (ValueError, TypeError):
            return str(val)
    return str(val) if not isinstance(val, (int, float, bool)) else val


def _write_header(ws, cols: list):
    ws.append(cols)
    for cell in ws[1]:
        cell.fill = FILL_HEADER
        cell.font = FONT_WHITE
        cell.border = THIN_BORDER
        cell.alignment = ALIGN_CENTER


def _write_rows(ws, data: list, cols: list):
    for doc in data:
        row_vals = [_fmt_value(doc.get(c), c) for c in cols]
        ws.append(row_vals)
        # Formato de fecha en celdas de fecha
        for ci, col_name in enumerate(cols, start=1):
            if col_name in DATE_COLS:
                cell = ws.cell(ws.max_row, ci)
                if isinstance(cell.value, datetime):
                    cell.number_format = "DD/MM/YYYY"


def _apply_cell_styles(ws, cols: list):
    """Aplica estilos celda a celda: fondo de fila, estado, responsable, crítico, info/review."""
    # Índices de columnas de interés (0-based para la fila)
    idx = {name: i for i, name in enumerate(cols)}

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.fill = FILL_ROW
            cell.border = THIN_BORDER
            cell.font = FONT_BLACK
            col_name = cols[cell.column - 1] if cell.column <= len(cols) else ""
            cell.alignment = ALIGN_CENTER if col_name in COLS_CENTRAR else ALIGN_VCE

        # Estado → color de celda
        if "Estado" in idx:
            estado_cell = row[idx["Estado"]]
            estado_val = str(estado_cell.value or "")
            if estado_val in ESTADO_COLORS:
                estado_cell.fill = PatternFill("solid", fgColor=ESTADO_COLORS[estado_val])
            if estado_val == "HOLD":
                estado_cell.font = Font(color="FF0000", bold=True)
            elif estado_val == "Aprobado":
                estado_cell.font = Font(bold=True)

        # Responsable → color de texto
        for resp_col in ("Responsable", "Repsonsable"):
            if resp_col in idx:
                resp_cell = row[idx[resp_col]]
                resp_val = str(resp_cell.value or "")
                if resp_val in RESPONSABLE_COLORS:
                    resp_cell.font = Font(color=RESPONSABLE_COLORS[resp_val], bold=True)

        # Crítico = Sí → rojo bold
        if "Crítico" in idx:
            crit_cell = row[idx["Crítico"]]
            if str(crit_cell.value or "").strip().lower() in ("sí", "si"):
                crit_cell.font = Font(color="FF0000", bold=True)

        # Info/Review: R → rojo bold, I → gris oscuro bold
        if "Info/Review" in idx:
            ir_cell = row[idx["Info/Review"]]
            if ir_cell.value == "R":
                ir_cell.font = Font(color="FF0000", bold=True)
            elif ir_cell.value == "I":
                ir_cell.font = Font(color="4D4D4D", bold=True)

        # Días Devolución > 15 → rojo bold
        if "Días Devolución" in idx:
            dd_cell = row[idx["Días Devolución"]]
            try:
                if dd_cell.value is not None and int(float(dd_cell.value)) > 15:
                    dd_cell.font = Font(color="FF0000", bold=True)
            except (ValueError, TypeError):
                pass


def _finalize_sheet(ws, cols: list):
    """Congela B2, autofilter y auto-ancho de columnas."""
    ws.freeze_panes = "B2"
    if ws.max_row >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"
    # Auto-ancho
    for col_idx, col_name in enumerate(cols, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(col_name)
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max_len + 2
