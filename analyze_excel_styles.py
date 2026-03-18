"""
Analyze Excel file styling for web replication.
Run with: docflow_env\Scripts\python.exe analyze_excel_styles.py
"""
import openpyxl
from openpyxl.utils import get_column_letter

path = r"U:\USUARIOS\jose.paredes\Desktop\DocuControl\monitoring_report_27-02-2026.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)

print("=" * 80)
print("SHEET NAMES:", wb.sheetnames)
print("=" * 80)

ws = wb.worksheets[0]
print(f"\nAnalyzing sheet: '{ws.title}'")
print(f"Dimensions: {ws.dimensions}")
print(f"Max row: {ws.max_row}, Max col: {ws.max_column}")

# Column widths
print("\n--- COLUMN WIDTHS ---")
for col_letter, dim in ws.column_dimensions.items():
    if dim.width:
        print(f"  Col {col_letter}: width={dim.width}, hidden={dim.hidden}")

# Row heights
print("\n--- ROW HEIGHTS (first 10) ---")
for r in range(1, min(11, ws.max_row + 1)):
    rd = ws.row_dimensions.get(r)
    if rd and rd.height:
        print(f"  Row {r}: height={rd.height}")

# Merged cells
print("\n--- MERGED CELLS ---")
for mc in ws.merged_cells.ranges:
    print(f"  {mc}")

def color_to_hex(color):
    if color is None:
        return "None"
    if color.type == "rgb" and color.rgb:
        return str(color.rgb)
    if color.type == "indexed":
        return f"indexed({color.indexed})"
    if color.type == "theme":
        return f"theme({color.theme}, tint={color.tint})"
    return str(color)

def describe_border_side(side):
    if side and side.style:
        c = color_to_hex(side.color) if side.color else "None"
        return f"{side.style} color={c}"
    return None

# Print header + first rows with full styling
max_col = min(ws.max_column, 30)
max_row_check = min(ws.max_row, 8)

for row_idx in range(1, max_row_check + 1):
    print(f"\n--- ROW {row_idx} ---")
    for col_idx in range(1, max_col + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        if cell.value is None and not (cell.fill and cell.fill.fill_type):
            continue
        col_l = get_column_letter(col_idx)
        print(f"  [{col_l}{row_idx}] value={repr(cell.value)}")

        # Fill
        fill = cell.fill
        if fill.fill_type:
            fg = color_to_hex(fill.fgColor)
            bg = color_to_hex(fill.bgColor)
            print(f"       fill: type={fill.fill_type}, fg={fg}, bg={bg}")

        # Font
        font = cell.font
        parts = []
        if font.name: parts.append(f"name={font.name}")
        if font.size: parts.append(f"size={font.size}")
        if font.bold: parts.append("BOLD")
        if font.italic: parts.append("ITALIC")
        if font.color: parts.append(f"color={color_to_hex(font.color)}")
        if parts:
            print(f"       font: {', '.join(parts)}")

        # Alignment
        al = cell.alignment
        if al.horizontal or al.vertical or al.wrap_text:
            print(f"       align: h={al.horizontal}, v={al.vertical}, wrap={al.wrap_text}")

        # Borders
        border = cell.border
        for side_name in ['left', 'right', 'top', 'bottom']:
            desc = describe_border_side(getattr(border, side_name))
            if desc:
                print(f"       border-{side_name}: {desc}")

        # Number format
        if cell.number_format and cell.number_format != 'General':
            print(f"       numfmt: {cell.number_format}")

# Conditional formatting
print("\n--- CONDITIONAL FORMATTING ---")
if ws.conditional_formatting:
    for cf in ws.conditional_formatting:
        print(f"  Range: {cf}")
        for rule in cf.rules:
            print(f"    Rule: type={rule.type}, operator={rule.operator}, formula={rule.formula}")
            if rule.dxf:
                dxf = rule.dxf
                if dxf.fill:
                    fg = color_to_hex(dxf.fill.fgColor) if dxf.fill.fgColor else 'None'
                    print(f"      fill: fg={fg}")
                if dxf.font:
                    fc = color_to_hex(dxf.font.color) if dxf.font.color else 'None'
                    print(f"      font: color={fc}, bold={dxf.font.bold}")
else:
    print("  None found")

# Other sheets summary
if len(wb.sheetnames) > 1:
    print("\n--- OTHER SHEETS SUMMARY ---")
    for sn in wb.sheetnames[1:]:
        s = wb[sn]
        print(f"  '{sn}': {s.max_row} rows x {s.max_column} cols")

wb.close()
print("\nDone.")
